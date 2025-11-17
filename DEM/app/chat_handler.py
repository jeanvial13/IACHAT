import os
import time
import json
from datetime import datetime, timedelta
from flask import (
    Flask,
    request,
    jsonify,
    render_template,
    redirect,
    url_for,
    session,
    send_file,
)
from werkzeug.utils import secure_filename
from openai import OpenAI

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from docx import Document  # for DEM Word uploads
except Exception:
    Document = None

try:
    from openpyxl import Workbook  # for Excel export
except Exception:
    Workbook = None

app = Flask(__name__, static_folder="static", template_folder="templates")
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "super-secret-key-change-me")

# Simple login using environment variables
APP_USER = os.environ.get("APP_USER")
APP_PASS = os.environ.get("APP_PASS")

UPLOAD_FOLDER = "uploads"
DEM_UPLOAD_FOLDER = os.path.join(UPLOAD_FOLDER, "dem_docs")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(DEM_UPLOAD_FOLDER, exist_ok=True)
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

client = OpenAI(api_key=os.environ.get("OPENAI_API_KEY"))
DEFAULT_MODEL = os.environ.get("OPENAI_MODEL", "gpt-4o-mini")

LOG_FILE = "chat_log.txt"
DEM_FILE = "dem_projects.json"


def _log(line: str) -> None:
    ts = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"[{ts}] {line}\n")


def extract_text(path: str) -> str:
    """Lee texto de TXT/MD o PDF (si pypdf está instalado)."""
    lower = path.lower()
    try:
        if lower.endswith(".txt") or lower.endswith(".md"):
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        if lower.endswith(".pdf") and PdfReader is not None:
            reader = PdfReader(path)
            parts = []
            for page in reader.pages:
                try:
                    parts.append(page.extract_text() or "")
                except Exception:
                    continue
            return "\n".join(parts)
    except Exception as e:
        _log(f"Error leyendo archivo {path}: {e}")
    return ""


def extract_docx_text(path: str) -> str:
    """Extrae texto de archivos .docx si python-docx está disponible."""
    if Document is None:
        _log("python-docx no está instalado; no puedo leer .docx")
        return ""
    try:
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        _log(f"Error leyendo DOCX {path}: {e}")
        return ""


# ==========================
#   DEMS JSON HELPERS
# ==========================


def _load_projects():
    if not os.path.exists(DEM_FILE):
        return []
    try:
        with open(DEM_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        _log(f"Error leyendo {DEM_FILE}: {e}")
        return []


def _save_projects(projects):
    try:
        with open(DEM_FILE, "w", encoding="utf-8") as f:
            json.dump(projects, f, ensure_ascii=False, indent=2)
    except Exception as e:
        _log(f"Error guardando {DEM_FILE}: {e}")


def _enrich_project(p):
    """Añade campos calculados (duración, nota reciente, stale)."""
    proj = dict(p)
    now = datetime.utcnow()

    # duración en días
    duration_days = None
    start_date_raw = proj.get("start_date")
    if start_date_raw:
        try:
            if len(start_date_raw) == 10:
                dt = datetime.strptime(start_date_raw, "%Y-%m-%d")
            else:
                dt = datetime.fromisoformat(start_date_raw)
            duration_days = (now - dt).days
        except Exception:
            duration_days = None
    proj["duration_days"] = duration_days

    # última nota
    notes = proj.get("notes") or []
    last_note_text = ""
    last_note_ts = None
    if notes:
        last = notes[-1]
        last_note_text = last.get("text", "")
        last_note_ts = last.get("created_at")
    proj["last_note"] = last_note_text

    # determinar fecha de última actualización
    updated_raw = proj.get("updated_at") or last_note_ts or start_date_raw
    stale = False
    if updated_raw:
        try:
            if len(updated_raw) == 10:
                udt = datetime.strptime(updated_raw, "%Y-%m-%d")
            else:
                udt = datetime.fromisoformat(updated_raw)
            stale = (now - udt) >= timedelta(days=5)
        except Exception:
            stale = False
    proj["stale"] = stale

    return proj


# ==========================
#   AUTH / LOGIN
# ==========================


@app.route("/login", methods=["GET", "POST"])
def login():
    error = None
    if request.method == "POST":
        username = (request.form.get("username") or "").strip()
        password = (request.form.get("password") or "").strip()

        if APP_USER and APP_PASS and username == APP_USER and password == APP_PASS:
            session["auth"] = True
            return redirect(url_for("home"))
        else:
            error = "Usuario o contraseña incorrectos."

    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ==========================
#   RUTAS PRINCIPALES
# ==========================


@app.route("/")
def home():
    # Protección SOLO para la página principal
    if not session.get("auth"):
        return redirect(url_for("login"))
    return render_template("index.html")


# DEM Manager NO protegido (como pediste)
@app.route("/dems")
def dems_page():
    return render_template("dem_manager.html")


# ==========================
#   CHAT IA BÁSICO
# ==========================


@app.route("/chat", methods=["POST"])
def chat():
    data = request.get_json() or {}
    message = data.get("message", "").strip()
    history = data.get("history", [])
    model = data.get("model") or DEFAULT_MODEL
    file_summaries = data.get("file_summaries", [])

    if not message:
        return jsonify({"error": "Mensaje vacío"}), 400

    messages = []
    for m in history:
        role = m.get("role")
        content = m.get("content")
        if role in ("user", "assistant") and isinstance(content, str):
            messages.append({"role": role, "content": content})

    user_content = message

    if file_summaries:
        user_content += "\n\n[Contexto de archivos adjuntos]\n"
        for fsum in file_summaries:
            fname = fsum.get("filename", "archivo")
            summ = fsum.get("summary", "")
            user_content += f"- {fname}: {summ}\n"

    messages.append({"role": "user", "content": user_content})

    try:
        completion = client.chat.completions.create(
            model=model,
            messages=messages,
        )
        reply = completion.choices[0].message.content
        _log(f"USER: {message}")
        _log(f"ASSISTANT: {reply}")
        return jsonify({"reply": reply})
    except Exception as e:
        _log(f"ERROR OpenAI: {e}")
        return jsonify({"error": f"Error llamando a OpenAI: {e}"}), 500


# ==========================
#   SUBIDA / RESUMEN DE ARCHIVOS
# ==========================


@app.route("/upload", methods=["POST"])
def upload():
    files = request.files.getlist("files")
    if not files:
        return jsonify({"error": "No se enviaron archivos"}), 400

    results = []

    for f in files:
        filename = secure_filename(f.filename or "archivo")
        save_name = f"{int(time.time())}_{filename}"
        path = os.path.join(app.config["UPLOAD_FOLDER"], save_name)
        _log(f"Guardando archivo en {path}")

        f.save(path)

        text = extract_text(path)
        if not text:
            results.append(
                {
                    "filename": filename,
                    "summary": "No pude leer este archivo (formato no soportado o vacío).",
                }
            )
            continue

        prompt = (
            "Eres un asistente que resume archivos para el usuario.\n"
            "Resume el contenido del archivo en español, máximo 120 palabras, "
            "de forma clara y con viñetas si es útil.\n\n"
            "Contenido del archivo:\n"
        )

        try:
            completion = client.chat.completions.create(
                model=DEFAULT_MODEL,
                messages=[
                    {"role": "system", "content": "Asistente para resumen de documentos."},
                    {"role": "user", "content": prompt + text[:8000]},
                ],
            )
            summary = completion.choices[0].message.content
        except Exception as e:
            summary = f"No pude resumir este archivo por un error con OpenAI: {e}"

        results.append({"filename": filename, "summary": summary})

    return jsonify({"files": results})


# ==========================
#   DEMS API
# ==========================


@app.route("/api/dems/projects", methods=["GET", "POST"])
def dem_projects():
    if request.method == "GET":
        projects = _load_projects()
        enriched = [_enrich_project(p) for p in projects]
        return jsonify({"projects": enriched})

    data = request.get_json() or {}
    now = datetime.utcnow().isoformat()

    project = {
        "id": int(time.time() * 1000),
        "name": data.get("name", "").strip(),
        "sponsor": data.get("sponsor", "").strip(),
        "requester": data.get("requester", "").strip(),
        "ba_owner": data.get("ba_owner", "").strip(),
        "title": data.get("title", "").strip(),
        "change_request": data.get("change_request", "").strip(),
        "cost_center": data.get("cost_center", "").strip(),
        "status": data.get("status", "Idea").strip(),
        "workflow_status": data.get("workflow_status", "Intake").strip(),
        "current_owner": data.get("current_owner", "").strip(),
        "start_date": data.get("start_date", now[:10]).strip(),
        "created_at": now,
        "updated_at": now,
        "notes": [],
        "doc_ai": "",
    }

    initial_note = data.get("initial_note", "").strip()
    if initial_note:
        project["notes"].append({"text": initial_note, "created_at": now})

    projects = _load_projects()
    projects.append(project)
    _save_projects(projects)

    return jsonify({"project": _enrich_project(project)}), 201


@app.route("/api/dems/projects/<int:proj_id>/note", methods=["POST"])
def dem_add_note(proj_id):
    data = request.get_json() or {}
    text = (data.get("text") or "").strip()
    if not text:
        return jsonify({"error": "Note text is required"}), 400

    projects = _load_projects()
    now = datetime.utcnow().isoformat()
    found = None

    for p in projects:
        if p.get("id") == proj_id:
            p.setdefault("notes", []).append({"text": text, "created_at": now})
            p["updated_at"] = now
            found = p
            break

    if not found:
        return jsonify({"error": "Project not found"}), 404

    _save_projects(projects)
    return jsonify({"project": _enrich_project(found)})


@app.route("/api/dems/projects/<int:proj_id>/attach", methods=["POST"])
def dem_attach_doc(proj_id):
    if "file" not in request.files:
        return jsonify({"error": "No file field 'file' in request"}), 400

    file = request.files["file"]
    if not file or file.filename == "":
        return jsonify({"error": "Empty file"}), 400

    filename = secure_filename(file.filename)
    save_name = f"{int(time.time())}_{filename}"
    path = os.path.join(DEM_UPLOAD_FOLDER, save_name)
    _log(f"[DEMS] Saving DOC for project {proj_id} at {path}")
    file.save(path)

    lower = filename.lower()
    if lower.endswith(".docx"):
        text = extract_docx_text(path)
    else:
        text = extract_text(path)

    if not text:
        return jsonify(
            {
                "error": "I could not read this document (unsupported or empty). "
                "Please upload a DOCX, TXT or PDF file."
            }
        ), 400

    user_prompt = (
        "You will receive the raw content of a project request document.\n\n"
        "1) Write a concise summary (max 200 words) of what is being requested.\n"
        "2) Provide an analysis of the project and propose possible solution options,\n"
        "   with a strong focus on SAP S/4HANA and other realistic technologies.\n\n"
        "Return the answer in English using two clearly labelled sections:\n"
        "=== SUMMARY ===\n"
        "...summary here...\n"
        "=== ANALYSIS & SOLUTIONS ===\n"
        "...analysis here...\n\n"
        "Document content follows:\n"
    )

    try:
        completion = client.chat.completions.create(
            model=DEFAULT_MODEL,
            messages=[
                {
                    "role": "system",
                    "content": "You are a senior IT Business Analyst and SAP S/4HANA solution architect.",
                },
                {
                    "role": "user",
                    "content": user_prompt + text[:8000],
                },
            ],
        )
        ai_text = completion.choices[0].message.content
    except Exception as e:
        _log(f"[DEMS] Error calling OpenAI for attach_doc: {e}")
        return jsonify({"error": f"Error calling OpenAI: {e}"}), 500

    projects = _load_projects()
    now = datetime.utcnow().isoformat()
    found = None

    for p in projects:
        if p.get("id") == proj_id:
            p["doc_ai"] = ai_text
            p["updated_at"] = now
            p.setdefault("notes", []).append(
                {"text": "AI analysis generated from attached document.", "created_at": now}
            )
            found = p
            break

    if not found:
        return jsonify({"error": "Project not found"}), 404

    _save_projects(projects)
    return jsonify({"project": _enrich_project(found), "doc_ai": ai_text})


@app.route("/api/dems/report", methods=["POST"])
def dem_report():
    """Genera un reporte en inglés de todos los proyectos usando la IA,
    con foco en el último estado e información más relevante."""
    projects = _load_projects()
    if not projects:
        return jsonify({"error": "There are no projects yet."}), 400

    now = datetime.utcnow().strftime("%Y-%m-%d")

    lines = []
    for p in projects:
        ep = _enrich_project(p)
        lines.append(f"Project ID: {ep.get('id')}")
        lines.append(f"DEM Name: {ep.get('name')}")
        lines.append(f"Sponsor: {ep.get('sponsor')}")
        lines.append(f"Requester: {ep.get('requester')}")
        lines.append(f"BA Owner: {ep.get('ba_owner')}")
        lines.append(f"Title: {ep.get('title')}")
        lines.append(f"Change requested: {ep.get('change_request')}")
        lines.append(f"Cost center: {ep.get('cost_center')}")
        lines.append(f"DEM status: {ep.get('status')}")
        lines.append(f"Workflow status: {ep.get('workflow_status')}")
        lines.append(f"Current task owner: {ep.get('current_owner')}")
        lines.append(f"Start date: {ep.get('start_date')}")
        lines.append(f"Duration (days): {ep.get('duration_days')}")
        lines.append(f"Is stale (>=5 days without updates): {ep.get('stale')}")
        lines.append(f"Last note: {ep.get('last_note')}")
        if ep.get("doc_ai"):
            lines.append("AI document insights are available for this project.")
        lines.append("---")

    context = "\n".join(lines)

    user_prompt = (
        "You are an experienced IT Portfolio Manager. "
        "Using the portfolio data below, write a detailed but concise report in English "
        "summarising the latest status of each project. For every project, include:\n"
        "- very short context (what it is about)\n"
        "- current status and workflow situation\n"
        "- most relevant recent information and last updates\n"
        "- key risks or blockers\n"
        "- clear next steps and owners\n\n"
        "Use headings with the DEM Name, bullet points and professional business language.\n"
        f"Today is {now}.\n\n"
        "Portfolio data:\n"
        + context
    )

    try:
        completion = client.chat.completions.create(
            model=DEFAULT_MODEL,
            messages=[
                {
                    "role": "system",
                    "content": "You are a senior IT Portfolio Manager writing executive status reports.",
                },
                {
                    "role": "user",
                    "content": user_prompt,
                },
            ],
        )
        report = completion.choices[0].message.content
    except Exception as e:
        _log(f"[DEMS] Error calling OpenAI for report: {e}")
        return jsonify({"error": f"Error calling OpenAI: {e}"}), 500

    return jsonify({"report": report})


@app.route("/api/dems/export", methods=["GET"])
def dem_export_excel():
    """Exporta todos los proyectos DEM a un archivo Excel."""
    if Workbook is None:
        return jsonify({"error": "openpyxl no está instalado en el entorno."}), 500

    projects = _load_projects()
    if not projects:
        return jsonify({"error": "There are no projects yet."}), 400

    wb = Workbook()
    ws = wb.active
    ws.title = "Projects"

    headers = [
        "id",
        "name",
        "sponsor",
        "requester",
        "ba_owner",
        "title",
        "change_request",
        "cost_center",
        "status",
        "workflow_status",
        "current_owner",
        "start_date",
        "duration_days",
        "stale",
        "last_note",
        "doc_ai_short",
    ]
    ws.append(headers)

    for p in projects:
        ep = _enrich_project(p)
        doc_ai = (ep.get("doc_ai") or "").replace("\n", " ")
        if len(doc_ai) > 300:
            doc_ai = doc_ai[:297] + "..."
        row = [
            ep.get("id"),
            ep.get("name"),
            ep.get("sponsor"),
            ep.get("requester"),
            ep.get("ba_owner"),
            ep.get("title"),
            ep.get("change_request"),
            ep.get("cost_center"),
            ep.get("status"),
            ep.get("workflow_status"),
            ep.get("current_owner"),
            ep.get("start_date"),
            ep.get("duration_days"),
            ep.get("stale"),
            ep.get("last_note"),
            doc_ai,
        ]
        ws.append(row)

    export_path = "dems_export.xlsx"
    wb.save(export_path)

    return send_file(
        export_path,
        as_attachment=True,
        download_name=f"dems_export_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx",
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8080)
